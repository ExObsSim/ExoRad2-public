import astropy.units as u
import openpyxl
import pytest
from astropy.table import Column, QTable

from exorad.models.targetlist import (
    CSVTargetList,
    CSVTargetListMRS,
    QTableTargetList,
    XLXSTargetList,
    compactString,
    stripUnitString,
)
from exorad.tasks.targetHandler import LoadTargetList

STAR_PLANET_CSV = """\
star name,star M [M_sun],star Teff [K],star R [R_sun],star D [pc],planet name,planet R [R_jup],planet M [M_jup]
HD 1,1.0,5800,1.0,12.0,HD 1 b,1.1,0.9
GJ 2,0.4,3400,0.4,5.0,GJ 2 b,0.3,0.05
"""

MRS_CSV = """\
star name,star mass [Ms],star temperature [K],star radius [Rs],star distance [pc],planet name,planet radius [Rj],planet mass [Mj],planet period [days]
HD 1,1.0,5800,1.0,12.0,HD 1 b,1.1,0.9,3.5
GJ 2,0.4,3400,0.4,5.0,GJ 2 b,0.3,0.05,1.2
"""


@pytest.fixture
def star_planet_csv(tmp_path):
    p = tmp_path / "targets.csv"
    p.write_text(STAR_PLANET_CSV)
    return str(p)


@pytest.fixture
def mrs_csv(tmp_path):
    p = tmp_path / "mrs.csv"
    p.write_text(MRS_CSV)
    return str(p)


@pytest.fixture
def target_xlsx(tmp_path):
    """Build a workbook with the layout XLXSTargetList expects.

    keys on sheet row 3, units on row 4, data from row 5; star in columns
    B:G, planet in columns I:S, N-to-observe in column U.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    star_keys = ["name", "M", "Teff", "R", "D", "magk"]
    star_units = ["", "M_sun", "K", "R_sun", "pc", ""]
    planet_keys = ["name", "R"] + [f"p{i}" for i in range(9)]

    for j, (k, unit) in enumerate(zip(star_keys, star_units, strict=True), start=2):
        ws.cell(row=3, column=j, value=k)
        ws.cell(row=3 + 1, column=j, value=unit)
    for j, k in enumerate(planet_keys, start=9):
        ws.cell(row=3, column=j, value=k)
    ws.cell(row=3, column=21, value="Nobs")

    rows = [
        ["HD 1", 1.0, 5800, 1.0, 12.0, 7.0, "HD 1 b", 1.1],
        ["GJ 2", 0.4, 3400, 0.4, 5.0, 6.0, "GJ 2 b", 0.3],
    ]
    for i, row in enumerate(rows, start=5):
        for j, val in enumerate(row[:6], start=2):
            ws.cell(row=i, column=j, value=val)
        ws.cell(row=i, column=9, value=row[6])
        ws.cell(row=i, column=10, value=row[7])
        ws.cell(row=i, column=21, value=1)

    path = tmp_path / "targets.xlsx"
    wb.save(path)
    return str(path)


def _raw_qtable():
    return QTable(
        [
            Column(["t1", "t2"], name="star name"),
            Column([1, 2] * u.M_sun, name="star M"),
            Column([5000, 6000] * u.K, name="star Teff"),
            Column([1, 2] * u.R_sun, name="star R"),
            Column([10, 12] * u.pc, name="star D"),
            Column(["t1 b", "t2 b"], name="planet name"),
            Column([1.0, 1.5] * u.R_jup, name="planet R"),
        ]
    )


class TestStringHelpers:
    def test_compact_string(self):
        assert compactString("HD 209458 b") == "hd209458b"

    def test_strip_unit_string(self):
        assert stripUnitString("[R_sun]") == "R_sun"


class TestCSVTargetList:
    def test_parses_star_and_planet(self, star_planet_csv):
        tl = CSVTargetList(star_planet_csv)
        assert len(tl.target) == 2
        first = tl.target[0]
        assert first.star.name == "HD 1"
        assert first.planet.name == "HD 1 b"
        assert first.star.M.unit == u.M_sun
        assert first.planet.R.unit == u.R_jup

    def test_search_target(self, star_planet_csv):
        tl = CSVTargetList(star_planet_csv)
        hits = tl.searchTarget("gj2")
        assert len(hits) == 1
        assert hits[0].star.name == "GJ 2"


class TestCSVTargetListMRS:
    def test_key_and_unit_translation(self, mrs_csv):
        tl = CSVTargetListMRS(mrs_csv)
        assert len(tl.target) == 2
        star = tl.target[0].star
        planet = tl.target[0].planet
        # "mass"/"radius"/... are renamed, "Ms"/"Rj"/"days" are mapped to astropy
        assert star.M.unit == u.M_sun
        assert star.Teff.unit == u.K
        assert planet.R.unit == u.R_jup
        assert planet.P.unit == u.day

    def test_check_units_and_keys_tables(self):
        assert CSVTargetListMRS.check_units("Me") == "M_earth"
        assert CSVTargetListMRS.check_units("unknown") == "unknown"
        assert CSVTargetListMRS.check_keys("semi-major axis") == "a"
        assert CSVTargetListMRS.check_keys("other") == "other"


class TestXLXSTargetList:
    def test_parses_workbook(self, target_xlsx):
        tl = XLXSTargetList(target_xlsx)
        assert len(tl.target) == 2
        assert tl.target[0].star.name == "HD 1"
        assert tl.target[0].star.M.unit == u.M_sun
        assert tl.target[0].planet.name == "HD 1 b"

    def test_dispatch_via_load_target_list(self, target_xlsx):
        out = LoadTargetList()(target_list=target_xlsx)
        assert isinstance(out, XLXSTargetList)


class TestQTableTargetList:
    def test_parses_quantities(self):
        tl = QTableTargetList(_raw_qtable())
        assert len(tl.target) == 2
        assert tl.target[0].star.M.unit == u.M_sun

    def test_rejects_non_quantity_columns(self):
        bad = _raw_qtable()
        bad["star D"] = Column([10, 12], name="star D")  # no unit
        with pytest.raises(TypeError):
            QTableTargetList(bad)


class TestLoadTargetListDispatch:
    def test_csv(self, star_planet_csv):
        out = LoadTargetList()(target_list=star_planet_csv)
        assert isinstance(out, CSVTargetList)
        assert len(out.target) == 2

    def test_qtable(self):
        out = LoadTargetList()(target_list=_raw_qtable())
        assert isinstance(out, QTableTargetList)

    def test_unknown_extension_raises(self, tmp_path):
        bad = tmp_path / "targets.txt"
        bad.write_text("nope")
        with pytest.raises(OSError):
            LoadTargetList()(target_list=str(bad))
